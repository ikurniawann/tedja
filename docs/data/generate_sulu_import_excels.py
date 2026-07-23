#!/usr/bin/env python3
"""Generate SULU Bandung import Excel files from docs/data/SULU-bdg.xlsx."""

from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs/data/SULU-bdg.xlsx"
OUT = ROOT / "docs/data/sulu-import"

WAREHOUSES = {
    "main storage": "WH-01",
    "hikiniku bar": "STALL-02",
    "noodles bar": "STALL-03",
    "onigiri corner": "STALL-04",
    "yokocho 1": "STALL-05",
    "yokocho 2": "STALL-06",
    "yokocho 3": "STALL-07",
    "yokocho 4": "STALL-08",
    "yokocho 5": "STALL-09",
    "yokocho 6": "STALL-10",
    "yokocho 7": "STALL-11",
    "yokocho 8": "STALL-14",
    "yokocho 9": "STALL-12",
    "yokocho 10": "STALL-13",
}

RM_HEADERS = [
    "kode",
    "nama",
    "kategori",
    "satuan_besar_kode",
    "satuan_kecil_kode",
    "konversi_factor",
    "stok_minimum",
    "stok_maximum",
    "shelf_life_days",
    "coa",
    "harga_beli",
    "opening_stock",
    "stall_code",
    "deskripsi",
    "status",
]

PRODUCT_HEADERS = [
    "kode",
    "nama",
    "stall_code",
    "kategori",
    "satuan_kode",
    "deskripsi",
    "harga_jual",
    "harga_modal",
    "markup_persen",
    "production_output_type",
    "status",
]

BOM_HEADERS = [
    "product_name",
    "product_type",
    "stall_code",
    "component_name",
    "component_type",
    "qty",
    "unit",
    "notes",
]


def norm(s: str | None) -> str:
    if s is None:
        return ""
    t = str(s).strip().lower()
    t = re.sub(r"\s+", " ", t)
    return t


def clean_name(s: str | None) -> str:
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s).strip())


def match_stall(station: str | None) -> tuple[str, str]:
    """Return (stall_code, note)."""
    if not station:
        return "", "UNMAPPED_STALL"
    key = norm(station)
    key = key.replace("yokocho", "yokocho ").replace("  ", " ").strip()
    # normalize "yokocho2" → "yokocho 2"
    key = re.sub(r"yokocho\s*(\d+)", r"yokocho \1", key)
    if key in WAREHOUSES:
        return WAREHOUSES[key], ""
    # fuzzy contains
    for name, code in WAREHOUSES.items():
        if name in key or key in name:
            return code, ""
    # Hikiniku / Onigiri / Noodles partial
    aliases = {
        "hikiniku": "STALL-02",
        "noodles": "STALL-03",
        "noodle": "STALL-03",
        "onigiri": "STALL-04",
        "ramen": "STALL-03",
        "udon": "STALL-03",
    }
    for a, code in aliases.items():
        if a in key:
            return code, f"STALL_ALIAS:{station}"
    return "", f"UNMAPPED_STALL:{station}"


def parse_unit(raw: str | None) -> tuple[str, str, float, str]:
    """Return besar, kecil, factor, note."""
    if not raw:
        return "PCS", "PCS", 1.0, "UNIT_DEFAULT_PCS"
    u = norm(raw)
    note = ""

    # bottle with volume
    m = re.search(r"(?:btl|botol).*?@?\s*([\d.,]+)\s*(ltr|l|ml|gr|g)?", u)
    if m or u.startswith("btl") or "botol" in u:
        vol = None
        unit_part = None
        m2 = re.search(r"([\d.,]+)\s*(ltr|liter|l|ml|gr|g)", u)
        if m2:
            vol = float(m2.group(1).replace(",", "."))
            unit_part = m2.group(2)
            if unit_part in ("ltr", "liter", "l"):
                return "BTL", "ML", vol * 1000, f"from:{raw}"
            if unit_part == "ml":
                return "BTL", "ML", vol, f"from:{raw}"
            if unit_part in ("gr", "g"):
                return "BTL", "GR", vol, f"from:{raw}"
        return "BTL", "BTL", 1.0, f"BTL_NO_VOLUME:{raw}"

    if u in ("kg",):
        return "KG", "GR", 1000.0, ""
    if u in ("gr", "g", "gram"):
        return "KG", "GR", 1000.0, "SOURCE_WAS_GR"
    if u in ("ltr", "liter", "l", "lt"):
        return "L", "ML", 1000.0, ""
    if u in ("ml",):
        return "L", "ML", 1000.0, "SOURCE_WAS_ML"
    if u in ("pack",):
        return "PACK", "PACK", 1.0, ""
    if u in ("pcs", "pc"):
        return "PCS", "PCS", 1.0, ""
    if u in ("buah",):
        return "PCS", "PCS", 1.0, "from:buah"
    if u in ("butir",):
        return "BUTIR", "BUTIR", 1.0, ""
    if u in ("ekor",):
        return "PCS", "PCS", 1.0, "from:ekor"
    if u in ("roll",):
        return "ROLL", "ROLL", 1.0, ""
    if "pinched" in u:
        return "GR", "GR", 1.0, f"APPROX_PINCH:{raw}"
    if u in ("sheets", "sheet", "lembar"):
        return "LBR", "LBR", 1.0, ""
    return "PCS", "PCS", 1.0, f"UNIT_FALLBACK:{raw}"


SEAFOOD_KW = (
    "ikan",
    "udang",
    "cumi",
    "squid",
    "salmon",
    "tuna",
    "unagi",
    "ebi",
    "seafood",
    "fish",
    "kakap",
    "cakalang",
    "mackarel",
    "mackerel",
    "shrimp",
    "nori",
    "katsuobushi",
    "tobiko",
    "mentaiko",
    "aonori",
    "nori",
)
DAIRY_KW = ("telur", "egg", "susu", "milk", "butter", "keju", "cheese", "mayo", "mayonnaise", "krim", "cream")
OIL_KW = ("minyak", "oil", "sesame oil", "olive")
SAUS_KW = ("sauce", "saus", "shoyu", "soy", "kecap", "gochujang", "mayo", "vinegar", "cuka", "tare", "oyster")
BUMBU_KW = (
    "garam",
    "gula",
    "lada",
    "pepper",
    "hondashi",
    "ajinomoto",
    "royco",
    "togarashi",
    "ichimi",
    "curry",
    "paprika powder",
    "jinten",
    "bumbu",
    "spice",
    "furikake",
    "baking powder",
    "vanilla",
)
BEKU_KW = ("frozen", "beku")
KEMASAN_KW = ("thinwall", "plastik", "foil", "baking paper", "wrapping", "box", "cup", "tissue")
NONPANG_KW = ("sponge", "sabun", "cuci", "sarung", "gloves", "cleaner")
BAKERY_KW = ("tepung", "flour", "chocolate", "cocoa", "ragi", "yeast", "breadcrumb", "panko")


def map_rm_category(market_cat: str | None, name: str) -> str:
    mc = norm(market_cat)
    n = norm(name)

    if mc == "fruit and vegetables":
        return "SAYUR"
    if mc == "kitchen supplies":
        if any(k in n for k in KEMASAN_KW):
            return "KEMASAN"
        return "NONPANG"
    if mc == "meat - poultry - fish":
        if any(k in n for k in SEAFOOD_KW):
            return "SEAFOOD"
        return "DAGING"
    if mc == "cold item":
        if any(k in n for k in BEKU_KW):
            return "BEKU"
        if any(k in n for k in DAIRY_KW):
            return "DAIRY"
        if any(k in n for k in SEAFOOD_KW):
            return "SEAFOOD"
        return "BEKU"
    if mc == "dry goods":
        if any(k in n for k in OIL_KW):
            return "OIL"
        if any(k in n for k in DAIRY_KW):
            return "DAIRY"
        if any(k in n for k in SAUS_KW):
            return "SAUS"
        if any(k in n for k in BUMBU_KW):
            return "BUMBU"
        if any(k in n for k in BAKERY_KW):
            return "BAKERY"
        if any(k in n for k in SEAFOOD_KW):
            return "SEAFOOD"
        return "KERING"

    # name-only fallback
    if any(k in n for k in SEAFOOD_KW):
        return "SEAFOOD"
    if any(k in n for k in DAIRY_KW):
        return "DAIRY"
    return "LAIN"


def map_product_category(menu_cat: str | None, is_wip: bool = False) -> str:
    if is_wip:
        return "OTHER"
    c = norm(menu_cat)
    if not c:
        return "OTHER"
    if "bakery" in c:
        return "BAKERY"
    if "rice" in c or "onigiri" in c:
        return "RICE"
    if "ramen" in c or "udon" in c or "noodle" in c:
        return "NOODLE"
    if "poridge" in c or "porridge" in c or "soup" in c:
        return "SOUP"
    if "sushi" in c or "sashimi" in c:
        return "SIDE"
    if "yakitori" in c or "dumpling" in c:
        return "SIDE"
    if "teppanyaki" in c or "agemono" in c or "yakimono" in c or "griiled" in c or "grilled" in c:
        return "MAIN"
    if "burger" in c or "sandwich" in c:
        return "MAIN"
    return "OTHER"


def opening_in_small(qty, besar: str, kecil: str, factor: float) -> float | str:
    if qty is None or qty == "":
        return ""
    try:
        q = float(qty)
    except (TypeError, ValueError):
        return ""
    if besar == kecil or factor == 1:
        return q
    # marketlist qty is usually in large unit (kg, ltr, pack)
    return q * factor


def write_sheet(path: Path, headers: list[str], rows: list[dict], extra_sheets: list[tuple[str, list[str], list[dict]]] | None = None):
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    bold = Font(bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.font = bold
    for r_idx, row in enumerate(rows, 2):
        for c_idx, h in enumerate(headers, 1):
            ws.cell(r_idx, c_idx, row.get(h, ""))
    if extra_sheets:
        for title, hdrs, erows in extra_sheets:
            w2 = wb.create_sheet(title[:31])
            for col, h in enumerate(hdrs, 1):
                w2.cell(1, col, h).font = bold
            for r_idx, row in enumerate(erows, 2):
                for c_idx, h in enumerate(hdrs, 1):
                    w2.cell(r_idx, c_idx, row.get(h, ""))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def load_marketlist(wb) -> list[dict]:
    items = []
    for sheet in ("Marketlist batch 1", "Marketlist batch 2"):
        ws = wb[sheet]
        for r in range(5, (ws.max_row or 0) + 1):
            name = clean_name(ws.cell(r, 3).value)
            if not name:
                continue
            cat = clean_name(ws.cell(r, 2).value) or None
            qty = ws.cell(r, 4).value
            unit = ws.cell(r, 5).value
            note = clean_name(ws.cell(r, 6).value)
            items.append(
                {
                    "source_sheet": sheet,
                    "market_cat": cat,
                    "nama": name,
                    "qty": qty,
                    "unit_raw": unit,
                    "note": note,
                }
            )
    return items


def dedupe_marketlist(items: list[dict]) -> list[dict]:
    by_key: dict[str, dict] = {}
    for it in items:
        key = norm(it["nama"])
        if key not in by_key:
            by_key[key] = dict(it)
            by_key[key]["sources"] = [it["source_sheet"]]
        else:
            prev = by_key[key]
            prev["sources"].append(it["source_sheet"])
            # prefer higher qty / keep first unit
            try:
                if float(it["qty"] or 0) > float(prev.get("qty") or 0):
                    prev["qty"] = it["qty"]
            except (TypeError, ValueError):
                pass
            if not prev.get("note") and it.get("note"):
                prev["note"] = it["note"]
            if not prev.get("market_cat") and it.get("market_cat"):
                prev["market_cat"] = it["market_cat"]
    return list(by_key.values())


def load_draft_menu(wb) -> list[dict]:
    ws = wb["Draft Menu"]
    menus = []
    for r in range(3, (ws.max_row or 0) + 1):
        name = clean_name(ws.cell(r, 2).value)
        if not name:
            continue
        menus.append(
            {
                "nama": name,
                "station": clean_name(ws.cell(r, 3).value),
                "menu_cat": clean_name(ws.cell(r, 4).value),
                "deskripsi": clean_name(ws.cell(r, 5).value),
                "batch": clean_name(ws.cell(r, 6).value),
            }
        )
    return menus


def parse_recipe_blocks(ws) -> list[dict]:
    blocks = []
    for r in range(1, (ws.max_row or 0) + 1):
        for c in (1, 8):
            if ws.cell(r, c).value != "NAME OF PRODUCT":
                continue
            product = clean_name(ws.cell(r, c + 2).value)
            category = None
            for rr in range(r, r + 5):
                if ws.cell(rr, c).value == "CATEGORY":
                    category = clean_name(ws.cell(rr, c + 2).value)
            header_row = None
            for rr in range(r, min(r + 10, (ws.max_row or 0) + 1)):
                if ws.cell(rr, c + 1).value == "Ingredients Item":
                    header_row = rr
                    break
            ingredients = []
            if header_row:
                for rr in range(header_row + 1, header_row + 40):
                    label = ws.cell(rr, c).value
                    if isinstance(label, str) and label.strip().lower().startswith("methode"):
                        break
                    item = clean_name(ws.cell(rr, c + 1).value)
                    qty = ws.cell(rr, c + 2).value
                    unit = ws.cell(rr, c + 3).value
                    if item:
                        ingredients.append({"name": item, "qty": qty, "unit": unit})
            if product:
                blocks.append(
                    {
                        "nama": product,
                        "category": category,
                        "ingredients": ingredients,
                        "row": r,
                        "col": c,
                    }
                )
    return blocks


# Heuristic WIP owner station from name / typical usage
WIP_OWNER_HINTS = {
    "cucumber pickles": "Yokocho 2",
    "marinasi beef": "Hikiniku Bar",
    "egg tartar": "Yokocho 2",
    "sulu spice": "Yokocho 2",
    "kimchi based": "Yokocho 7",
    "okonomiyaki sauce": "Yokocho 3",
    "soy chicken": "Yokocho 7",
    "namban sauce": "Yokocho 7",
    "tare sauce": "Yokocho 2",
    "furikake": "Yokocho 2",
    "adonan gyoza": "Yokocho 4",
    "mentaiko mayo": "Yokocho 4",
    "spicy mayo": "Yokocho 4",
    "tteokbokki sauce": "Yokocho 4",
    "yakiniku sauce": "Hikiniku Bar",
    "ebi shinjo": "Yokocho 7",
}

# Recipe ingredient label → marketlist / canonical RM name
RM_ALIASES = {
    "kyuri/japanese cucumber": "Kyuri / Timun Jepang",
    "kyuri": "Kyuri / Timun Jepang",
    "japanese cucumber": "Kyuri / Timun Jepang",
    "sesame oil": "Minyak Wijen",
    "minyak sesame": "Minyak Wijen",
    "minyak wijen": "Minyak Wijen",
    "air": "Air",
    "garam": "Garam Refina",
    "garam (taburkan pada timun saja)": "Garam Refina",
    "gula": "Gula Putih",
    "gula putih": "Gula Putih",
    "sugar/gula putih": "Gula Putih",
    "brown sugar": "Gula Merah",
    "gula merah": "Gula Merah",
    "shoyu": "Shoyu Kikkoman Halal",
    "shoyu kikkoman": "Shoyu Kikkoman Halal",
    "soy sauce kikkoman": "Shoyu Kikkoman Halal",
    "soy sauce kikkoman halal": "Shoyu Kikkoman Halal",
    "kikkoman soy sauce halal": "Shoyu Kikkoman Halal",
    "soy kikkoman": "Shoyu Kikkoman Halal",
    "hondashi": "Hondashi",
    "bawang bombay slice": "Bawang Bombay",
    "bombay": "Bawang Bombay",
    "bombay putih / onion chopped": "Bawang Bombay",
    "bawang bombay blender/parut": "Bawang Bombay",
    "beef shortplate slice": "Shortplates Beef slice (yakiniku size / 2mm)",
    "telur rebus dan chopped": "Telur Ayam Curah",
    "telur": "Telur Ayam Curah",
    "kewpie mayonnaise": "Kewpie Mayonaise",
    "mayonaise": "Kewpie Mayonaise",
    "mayonnaise": "Kewpie Mayonaise",
    "cucumber pickles chopped": "Cucumber Pickles",  # WIP
    "black pepper": "Lada Putih",
    "lada hitam bubuk": "Lada Putih",
    "lada putih bubuk": "Lada Putih",
    "paprika powder": "Paprika Merah",
    "togarashi": "Togarashi / Japanese Chili Powder",
    "togarashi / chili powder": "Togarashi / Japanese Chili Powder",
    "ichimi togarashi": "Togarashi / Japanese Chili Powder",
    "jinten powder": "Jinten Powder",
    "gochujang paste": "Gochujang Paste",
    "jahe": "Jahe",
    "jahe parut": "Jahe",
    "bawang putih kupas": "Bawang putih",
    "bawang putih": "Bawang putih",
    "bawang putih parut": "Bawang putih",
    "bawang putih blender/parut": "Bawang putih",
    "cabe kering": "Chili Flakes",
    "dry chili": "Chili Flakes",
    "ajinomoto": "Ajinomoto",
    "royco": "Royco ayam",
    "rice vinegar": "Rice Vinegar",
    "sesame seed": "Biji Wijen Putih",
    "white sesame": "Biji Wijen Putih",
    "wijen putih": "Biji Wijen Putih",
    "daun kucai slice": "Daun Kucai",
    "sauce tomat": "Sauce Tomat",
    "madu": "Madu",
    "oyster sauce": "Oyster Sauce",
    "curry powder": "Curry Powder",
    "paha ayam fillet tanpa kulit": "Paha Ayam Fillet",
    "tulang ceker ayam": "Ceker Ayam",
    "daun bawang": "Daun Bawang Besar",
    "daun bawang slice": "Daun Bawang Besar",
    "wortel": "Wortel Brastagi",
    "yakinori sheets": "Nori Sheets / yakinori",
    "katsuobushi": "Katsuobushi",
    "ayam giling": "Ayam Giling",
    "sawi putih chopped": "Sawi Putih",
    "campignon chopped": "Jamur Campignon",
    "mentaiko paste": "Mentaiko Paste",
    "chili sauce": "Chili sauce",
    "butter unsalted": "Unsalted Butter",
    "fresh milk": "Fresh Milk white",
    "apel fuji blender/parut (kupas)": "Apel Fuji",
    "lemon juice/peras": "Lemon Import",
    "dry shitake": "Jamur Shitake Kering",
    "udang vaname": "Udang Vaname",
    "cumi": "Cumi",
    "tepung tapioka": "Tepung Tapioka",
    "tepung roti": "Tepung Roti",
    "aonori": "Nori Flakes (Aonori)",
    "es batu": "Es Batu",
    "tapioka charcoal": "Tepung Tapioka",
    "egg tartar mayo": "Egg tartar",  # WIP
    "namban sauce": "Namban Sauce",  # WIP
    "furikake": "Furikake",  # WIP
    "soy chicken": "Soy Chicken",  # WIP
}


def make_rm_row(
    nama: str,
    kategori: str,
    unit_raw,
    qty,
    cat_counters: dict[str, int],
    deskripsi: str,
) -> dict:
    cat_counters[kategori] += 1
    kode = f"BB-{kategori}-{cat_counters[kategori]:03d}"
    besar, kecil, factor, unit_note = parse_unit(unit_raw)
    coa = "ASSET" if kategori in ("NONPANG", "KEMASAN", "BAKAR") else "PRODUCTION"
    parts = [p for p in [deskripsi, unit_note] if p]
    opening = opening_in_small(qty, besar, kecil, factor)
    return {
        "kode": kode,
        "nama": nama,
        "kategori": kategori,
        "satuan_besar_kode": besar,
        "satuan_kecil_kode": kecil,
        "konversi_factor": factor,
        "stok_minimum": "",
        "stok_maximum": "",
        "shelf_life_days": "",
        "coa": coa,
        "harga_beli": "",
        "opening_stock": opening,
        "stall_code": "WH-01",
        "deskripsi": " | ".join(parts),
        "status": "active",
    }


def normalize_bom_unit(unit_raw) -> str:
    bom_unit = str(unit_raw or "").strip().upper()
    if bom_unit in ("GR", "G", "GRAM"):
        return "GR"
    if bom_unit in ("ML",):
        return "ML"
    if bom_unit in ("KG",):
        return "KG"
    if "pinch" in norm(unit_raw):
        return "GR"
    if bom_unit in ("BUAH", "PCS", "PC"):
        return "PCS"
    if "sheet" in norm(unit_raw):
        return "LBR"
    return bom_unit or "PCS"


def main():
    wb = load_workbook(SRC, data_only=True)

    # --- Raw materials from marketlist ---
    market = dedupe_marketlist(load_marketlist(wb))
    rm_rows: list[dict] = []
    rm_by_norm: dict[str, str] = {}  # norm → canonical nama
    cat_counters: dict[str, int] = defaultdict(int)

    for it in sorted(market, key=lambda x: (x.get("market_cat") or "", x["nama"].lower())):
        kategori = map_rm_category(it.get("market_cat"), it["nama"])
        desc_parts = []
        if it.get("market_cat"):
            desc_parts.append(f"Marketlist: {it['market_cat']}")
        if it.get("note"):
            desc_parts.append(it["note"])
        if len(it.get("sources", [])) > 1:
            desc_parts.append("deduped:" + "+".join(sorted(set(it["sources"]))))
        row = make_rm_row(
            it["nama"], kategori, it.get("unit_raw"), it.get("qty"), cat_counters, " | ".join(desc_parts)
        )
        rm_rows.append(row)
        rm_by_norm[norm(it["nama"])] = it["nama"]

    # --- Products FG ---
    menus = load_draft_menu(wb)
    product_rows: list[dict] = []
    product_index: dict[str, dict] = {}
    for m in menus:
        stall, stall_note = match_stall(m.get("station"))
        if not stall:
            stall, stall_note = match_stall(m.get("menu_cat"))
        notes = []
        if stall_note:
            notes.append(stall_note)
        if not stall:
            stall = "STALL-05"
            notes.append("STALL_FALLBACK_YOKOCHO1")
        desc = m.get("deskripsi") or ""
        if notes:
            desc = (desc + " | " if desc else "") + "; ".join(notes)
        row = {
            "kode": "",
            "nama": m["nama"],
            "stall_code": stall,
            "kategori": map_product_category(m.get("menu_cat")),
            "satuan_kode": "PORSI",
            "deskripsi": desc,
            "harga_jual": "",
            "harga_modal": "",
            "markup_persen": "",
            "production_output_type": "FINISHED_GOOD",
            "status": "active",
        }
        product_rows.append(row)
        product_index[norm(m["nama"])] = row

    # --- WIP from SUB-RECIPE ---
    sub = parse_recipe_blocks(wb["SUB-RECIPE CARD"])
    sales = parse_recipe_blocks(wb["ITEM SALES RECIPE"])

    wip_by_norm: dict[str, str] = {}
    for block in sub:
        wip_by_norm[norm(block["nama"])] = block["nama"]
        hint = WIP_OWNER_HINTS.get(norm(block["nama"]))
        stall, stall_note = match_stall(hint) if hint else ("", "WIP_NO_HINT")
        if not stall:
            stall = "STALL-05"
            stall_note = (stall_note + ";" if stall_note else "") + "STALL_FALLBACK_YOKOCHO1"
        desc = "Sub-recipe WIP from SUB-RECIPE CARD"
        if stall_note:
            desc += f" | {stall_note}"
        row = {
            "kode": "",
            "nama": block["nama"],
            "stall_code": stall,
            "kategori": "OTHER",
            "satuan_kode": "PORSI",
            "deskripsi": desc,
            "harga_jual": "",
            "harga_modal": "",
            "markup_persen": "",
            "production_output_type": "WIP",
            "status": "active",
        }
        if norm(block["nama"]) in product_index:
            product_index[norm(block["nama"])]["deskripsi"] = (
                (product_index[norm(block["nama"])].get("deskripsi") or "") + " | ALSO_IN_SUBRECIPE"
            )
        else:
            product_rows.append(row)
            product_index[norm(block["nama"])] = row

    for block in sales:
        key = norm(block["nama"])
        # match draft menu ignoring case
        matched = None
        for pk, prow in product_index.items():
            if pk == key or pk.replace(" ", "") == key.replace(" ", ""):
                matched = prow
                break
            if "namban" in key and "namban" in pk and "burger" not in pk:
                matched = prow
                break
            if "torikaraage" in key and "karaage" in pk:
                matched = prow
                break
        if matched:
            product_index[key] = matched
            continue
        stall = "STALL-11"
        row = {
            "kode": "",
            "nama": block["nama"].title() if block["nama"].isupper() else block["nama"],
            "stall_code": stall,
            "kategori": "MAIN",
            "satuan_kode": "PORSI",
            "deskripsi": "From ITEM SALES RECIPE",
            "harga_jual": "",
            "harga_modal": "",
            "markup_persen": "",
            "production_output_type": "FINISHED_GOOD",
            "status": "active",
        }
        product_rows.append(row)
        product_index[key] = row
        product_index[norm(row["nama"])] = row

    # Collect recipe ingredients → ensure RM/WIP coverage via aliases + auto-add
    all_ings: list[tuple[str, object]] = []
    for block in sub + sales:
        for ing in block["ingredients"]:
            all_ings.append((ing["name"], ing.get("unit")))

    def canonical_from_alias(name: str) -> str | None:
        n = norm(name)
        n_base = re.sub(r"\s+", " ", re.sub(r"\s*\(.*?\)\s*", " ", n)).strip()
        if n in RM_ALIASES:
            return RM_ALIASES[n]
        if n_base in RM_ALIASES:
            return RM_ALIASES[n_base]
        return None

    for ing_name, unit_raw in all_ings:
        alias = canonical_from_alias(ing_name)
        target = alias or clean_name(ing_name)
        tn = norm(target)
        if tn in wip_by_norm:
            continue
        if tn in rm_by_norm:
            continue
        # fuzzy against existing RM
        found = None
        for rn, display in rm_by_norm.items():
            if tn in rn or rn in tn:
                found = display
                break
        if found:
            rm_by_norm[tn] = found
            continue
        # auto-add missing RM from recipe
        kategori = map_rm_category(None, target)
        # refine water / honey etc.
        if tn == "air":
            kategori = "MINUMAN"
        elif "madu" in tn:
            kategori = "SAUS"
        elif "ajinomoto" in tn:
            kategori = "BUMBU"
        elif "jinten" in tn:
            kategori = "BUMBU"
        elif "mentaiko" in tn:
            kategori = "SEAFOOD"
        elif "udang" in tn or "cumi" in tn:
            kategori = "SEAFOOD"
        elif "chili sauce" in tn:
            kategori = "SAUS"
        elif "tepung roti" in tn:
            kategori = "BAKERY"
        elif "es batu" in tn:
            kategori = "MINUMAN"
        row = make_rm_row(
            target,
            kategori,
            unit_raw,
            None,
            cat_counters,
            f"Auto-added from recipe ingredient (source label: {ing_name})",
        )
        rm_rows.append(row)
        rm_by_norm[tn] = target
        rm_by_norm[norm(ing_name)] = target

    # --- Product BOM ---
    def resolve_component(name: str) -> tuple[str, str, str]:
        n = norm(name)
        n_base = re.sub(r"\s+", " ", re.sub(r"\s*\(.*?\)\s*", " ", n)).strip()
        alias = canonical_from_alias(name)
        candidates = [n, n_base, norm(alias) if alias else ""]

        for c in candidates:
            if c and c in wip_by_norm:
                return "WIP", wip_by_norm[c], "ALIAS" if alias and norm(alias) == c else ""
        if alias and norm(alias) in wip_by_norm:
            return "WIP", wip_by_norm[norm(alias)], "ALIAS"

        for c in candidates:
            if c and c in rm_by_norm:
                note = "ALIAS" if alias and rm_by_norm[c] == alias else ""
                return "RAW", rm_by_norm[c], note

        # fuzzy WIP
        for wn, display in wip_by_norm.items():
            if n_base and (n_base in wn or wn in n_base):
                return "WIP", display, f"FUZZY_WIP:{display}"
        # fuzzy RM
        for rn, display in rm_by_norm.items():
            if n_base and (n_base in rn or rn in n_base):
                return "RAW", display, f"FUZZY_RM:{display}"
        return "RAW", alias or clean_name(name), "UNMAPPED_COMPONENT"

    bom_rows: list[dict] = []
    for block in sub:
        prow = product_index.get(norm(block["nama"]), {})
        for ing in block["ingredients"]:
            ctype, cname, cnote = resolve_component(ing["name"])
            bom_rows.append(
                {
                    "product_name": block["nama"],
                    "product_type": "WIP",
                    "stall_code": prow.get("stall_code", ""),
                    "component_name": cname,
                    "component_type": ctype,
                    "qty": ing.get("qty") if ing.get("qty") is not None else "",
                    "unit": normalize_bom_unit(ing.get("unit")),
                    "notes": cnote,
                }
            )

    for block in sales:
        key = norm(block["nama"])
        prow = product_index.get(key) or {}
        pname = prow.get("nama") or block["nama"]
        for ing in block["ingredients"]:
            ctype, cname, cnote = resolve_component(ing["name"])
            bom_rows.append(
                {
                    "product_name": pname,
                    "product_type": "FINISHED_GOOD",
                    "stall_code": prow.get("stall_code", ""),
                    "component_name": cname,
                    "component_type": ctype,
                    "qty": ing.get("qty") if ing.get("qty") is not None else "",
                    "unit": normalize_bom_unit(ing.get("unit")),
                    "notes": cnote,
                }
            )

    OUT.mkdir(parents=True, exist_ok=True)
    write_sheet(OUT / "01-raw-materials.xlsx", RM_HEADERS, rm_rows)
    write_sheet(OUT / "02-products.xlsx", PRODUCT_HEADERS, product_rows)
    write_sheet(
        OUT / "03-product-bom.xlsx",
        BOM_HEADERS,
        bom_rows,
        extra_sheets=[
            (
                "unmapped",
                BOM_HEADERS,
                [r for r in bom_rows if "UNMAPPED" in (r.get("notes") or "")],
            )
        ],
    )
    write_sheet(
        OUT / "04-raw-material-bom.xlsx",
        ["note"],
        [
            {
                "note": "Tidak dipakai. Sub-recipe SULU dipetakan sebagai Product WIP + Product BOM (lihat 02-products.xlsx dan 03-product-bom.xlsx)."
            }
        ],
    )

    unmapped = [r for r in bom_rows if "UNMAPPED" in (r.get("notes") or "")]
    print(f"Raw materials: {len(rm_rows)}")
    print(
        f"Products: {len(product_rows)} "
        f"(FG={sum(1 for r in product_rows if r['production_output_type']=='FINISHED_GOOD')}, "
        f"WIP={sum(1 for r in product_rows if r['production_output_type']=='WIP')})"
    )
    print(f"Product BOM lines: {len(bom_rows)} (unmapped components: {len(unmapped)})")
    print(f"Wrote → {OUT}")


if __name__ == "__main__":
    main()
